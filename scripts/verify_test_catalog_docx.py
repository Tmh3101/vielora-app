#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Verify the generated test catalog DOCX — corporate layout checks."""
from docx import Document
from docx.shared import Cm

path = "/home/hieutm/Work/Titops/Vielora/vielora/docs/manual-testing/Vielora_Manual_Test_Catalog.docx"
doc = Document(path)

# --- Layout checks ---
sec = doc.sections[0]
is_a4 = abs(sec.page_width.cm - 21.0) < 0.5 and abs(sec.page_height.cm - 29.7) < 0.5
margin_ok = all(abs(m.cm - 2.5) < 0.2 for m in
                [sec.top_margin, sec.bottom_margin, sec.left_margin, sec.right_margin])
diff_first_page = sec.different_first_page_header_footer

hdr_text = " | ".join(p.text for p in sec.header.paragraphs if p.text.strip())
ftr_text = " | ".join(p.text for p in sec.footer.paragraphs if p.text.strip())
hdr_has_title = "Vielora" in hdr_text and "4.1" in hdr_text
ftr_has_field = "PAGE" in sec.footer._element.xml and "NUMPAGES" in sec.footer._element.xml
ftr_has_conf = "nội bộ" in ftr_text

# --- Heading hierarchy ---
h1 = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
h2 = [p.text for p in doc.paragraphs if p.style.name == "Heading 2"]
groups_h1 = [t for t in h1 if t.startswith("NHÓM") or t.startswith("PHỤ LỤC") or t in
             ("MÔI TRƯỜNG KIỂM THỬ", "TÀI KHOẢN GỢI Ý", "LỊCH SỬ THAY ĐỔI", "MỤC LỤC")]
title_style = [p.text for p in doc.paragraphs if p.style.name == "Title"]

# --- TOC field ---
toc_found = "TOC \\o" in doc._element.xml or "TOC \\\\o" in doc._element.xml

# --- Literal > bug ---
literal_gt = [p.text for p in doc.paragraphs if p.text.strip().startswith(">")]

# --- Page breaks before H1 ---
breaks = sum(1 for p in doc.paragraphs if 'type="page"' in p._element.xml)
page_breaks_ok = breaks >= 14  # ~16 H1 sections, break trước hầu hết

# --- Content checks ---
text = "\n".join(p.text for p in doc.paragraphs)
for t in doc.tables:
    for row in t.rows:
        for cell in row.cells:
            text += "\n" + cell.text
diacritics = "àáảãạăâđèéêìíòóôõùúýỳỷỹỵửữừứơớờởỡ"
viet_count = sum(1 for ch in text if ch in diacritics)

# --- Numbered list restart check (mỗi list bắt đầu từ 1) ---
import re as _re
num_lists = []
cur = []
for p in doc.paragraphs:
    m = _re.match(r"^\s*(\d+)\.\s", p.text)
    if m:
        cur.append(int(m.group(1)))
    else:
        if cur:
            num_lists.append(cur)
            cur = []
if cur:
    num_lists.append(cur)
lists_ok = all(lst == list(range(1, len(lst) + 1)) for lst in num_lists) and len(num_lists) >= 5

checks = {
    "A4 page": is_a4,
    "Margin 2.5cm": margin_ok,
    "First page khac (cover sach)": diff_first_page,
    "Header co ten + version": hdr_has_title,
    "Footer co Trang X/Y (fields)": ftr_has_field,
    "Footer co bao mat": ftr_has_conf,
    "TOC field": toc_found,
    "Khong con literal '>'": len(literal_gt) == 0,
    "H1 >= 16 (sections+groups)": len(h1) >= 16,
    "H2 >= 55 (features)": len(h2) >= 55,
    "Group la H1 khong phai Title": len(title_style) == 0,
    "Page breaks giua cac nhom": page_breaks_ok,
    "Cover co ma tai lieu": "VL-DOC-TEST-001" in text,
    "Version 4.1": "4.1" in text,
    "Bang kiem soat 5 cot": any(len(t.rows[0].cells) == 5 for t in doc.tables),
    "F01-F59 day du": "F01" in text and "F59" in text,
    "Co lap du lieu": "cô lập dữ liệu" in text.lower(),
    "Placeholder dien": "[ĐIỀN" in text,
    "Khong con jargon": "BullMQ" not in text and "/api/" not in text,
    "Tieng Viet day du": viet_count > 1500,
    "List so bat dau tu 1": lists_ok,
    "Anh minh hoa placeholder >= 55": text.count("Ảnh minh họa") >= 55,
}

print(f"A4={is_a4} | Margins={margin_ok} | H1={len(h1)} | H2={len(h2)} | Tables={len(doc.tables)}")
print(f"Header: [{hdr_text}]")
print(f"Footer: [{ftr_text}]")
print(f"Vietnamese chars: {viet_count} | Page breaks: {breaks} | literal '>': {len(literal_gt)}")
print()
all_ok = True
for k, v in checks.items():
    print(f"  {'OK  ' if v else 'MISS'} {k}")
    all_ok = all_ok and v

print("\nVERIFY:", "PASS" if all_ok else "FAIL")
