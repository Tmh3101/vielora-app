#!/usr/bin/env python3
# Highlight (shading) the paragraphs/headings that were updated, with TWO
# distinct colors to separate edit rounds:
#   - YELLOW  : subscription-flow sync round (v4.2-4.3): NHÓM 7 F38-F43.1,
#               F44-F46; NHÓM 9 F49, F51.
#   - GREEN   : group-chat round (v4.5): F66 (Pin knowledge), F69 (Daily summary).
# Run AFTER generate_test_catalog_docx.py (which produces a clean, unshaded docx).

import sys
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

YELLOW = "FFF2CC"  # light yellow  — round 1 (subscription)
GREEN = "E2EFDA"   # light green   — round 2 (group chat)

# Markers per round: any heading whose text starts with one of these.
YELLOW_MARKERS = [
    "F38.", "F39.", "F40.", "F41.", "F42.", "F43.",
    "F43.1",
    "F44.", "F45.", "F46.",
    "F49.", "F51.",
]
GREEN_MARKERS = [
    "F66.", "F69.",
]


def shade_paragraph(p, color):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    pPr.append(shd)


def markers_of(p, markers):
    t = p.text.strip()
    for m in markers:
        if t.startswith(m) or t == m.strip():
            return True
    return False


def region_shade(doc, markers, color):
    """Shade a heading and everything below it until the next feature/group
    heading that is NOT in `markers`."""
    target_headings = set()
    for p in doc.paragraphs:
        if markers_of(p, markers):
            target_headings.add(p.text.strip())

    shaded = 0
    shading_on = False
    for p in doc.paragraphs:
        t = p.text.strip()
        is_marker_heading = t in target_headings
        looks_like_f = (len(t) >= 3 and t[0] == "F" and t[1].isdigit() and
                        (t[2] == "." or (t[2].isdigit() and t[3] == ".")))
        is_group = (t.startswith("NHÓM") or t.startswith("PHỤ LỤC")
                    or t.startswith("MÔI") or t.startswith("TÀI")
                    or t.startswith("LỊCH") or t.startswith("MỤC"))

        if is_marker_heading:
            shading_on = True
        elif looks_like_f or is_group:
            if not any(t.startswith(m.rstrip(".")) for m in markers):
                shading_on = False

        if shading_on or is_marker_heading:
            shade_paragraph(p, color)
            shaded += 1

    # Also shade table rows whose first cell starts with a marker
    for tbl in doc.tables:
        for row in tbl.rows:
            cells = [c.text.strip() for c in row.cells]
            head = cells[0] if cells else ""
            if any(head.startswith(m.rstrip(".")) for m in markers):
                for cell in row.cells:
                    shade_paragraph(cell.paragraphs[0], color)
                    shaded += 1

    return shaded


def main():
    if len(sys.argv) != 2:
        print("Usage: highlight_updates_docx.py <file.docx>")
        sys.exit(1)
    path = sys.argv[1]
    doc = Document(path)

    y = region_shade(doc, YELLOW_MARKERS, YELLOW)
    g = region_shade(doc, GREEN_MARKERS, GREEN)

    doc.save(path)
    print(f"Highlighted {y} (yellow, subscription round) + {g} "
          f"(green, group-chat round) paragraphs/cells in {path}")


if __name__ == "__main__":
    main()
